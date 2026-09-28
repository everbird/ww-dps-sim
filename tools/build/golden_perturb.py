"""golden 扰动对拍（TD-03 §10.3）：开发期工具，不进 CI。

xlsx 的缓存值只用了一套配置，大多数乘区是 0，只能证明它们的中性值。这里把标准答案公式引用到的「伤害配置」数值格
（R1791 起的分区网格与汇总格）、目标防御与七项抗性随机改写：用解析器按原数组公式算出"期望值"，再按 TD-03 §10.2
拆乘区、按 §3–§7 的公式重算，逐格比较。改公式、改拆分规则、换 xlsx 版本后跑一次。

用法：pnpm golden:perturb [-- --rounds 20 --seed 1]（即 python3 tools/build/golden_perturb.py [xlsx] …）

已知差异单独计数、不算失败：防御分母 ≤ 0（减防大到有效防御 ≤ −(800 + 8·Lv)）时 xlsx 算出负的防御系数，TD-03 取上限 2
（§1.3）。物理、谐度破坏与响应"减防后先取整"的写法（§3.2、Q1）在这里不会出现差异：拆分时把 FLOOR(…) 整体记为目标防御。
"""
from __future__ import annotations

import argparse
import random
import sys
import time
import warnings
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).parent))
import golden  # noqa: E402
from xlformula import Evaluator, XlError, col_letter, refs_in, split_coord  # noqa: E402

ROOT = Path(__file__).resolve().parents[2]
CALC, CFG = golden.CALC, golden.CFG
TARGET_CELLS = ['U3'] + [f'{col_letter(c)}3' for c in range(22, 29)]   # 伤害计算 U3 目标防御，V3–AB3 七项抗性


def inputs_of(ev: Evaluator, cells: list[tuple[str, str, str]]) -> tuple[set, set]:
    """(要改写的格, 要按公式现算的格)：标准答案公式引用到的「伤害配置」R1791 起的数值格；
    其中预先算好的因子格（R1792–R1815 防御 / 抗性系数、R2139–R2142 霜冻 1002 类 MAX(1 + x, 0)）改为现算，
    它们公式里引用的格也一并改写"""
    cfg_f = ev.formulas.get(CFG, {})
    cfg_v = ev.values.get(CFG, {})
    inputs: set[tuple[str, str]] = set()
    live: set[tuple[str, str]] = set()
    todo = [(CALC, ev.formulas[CALC][c]) for c, _, _ in cells]
    while todo:
        here, src = todo.pop()
        for ref in refs_in(ev.parsed(src)):
            sheet = ref.sheet or here
            if sheet != CFG or ref.r1 is None or ref.r1 < 1791:
                continue
            for r in range(ref.r1, ref.r2 + 1):
                for c in range(ref.c1, ref.c2 + 1):
                    key = (CFG, f'{col_letter(c)}{r}')
                    if (1792 <= r <= 1815 or 2139 <= r <= 2142) and isinstance(cfg_f.get(key[1]), str) \
                            and cfg_f[key[1]].startswith('='):
                        if key not in live:
                            live.add(key)
                            todo.append((CFG, cfg_f[key[1]]))
                        continue
                    v = cfg_v.get(key[1])
                    if v is None or (isinstance(v, (int, float)) and not isinstance(v, bool)):
                        inputs.add(key)
    inputs |= {(CALC, c) for c in TARGET_CELLS}
    return inputs, live


def draw(rng: random.Random, key: tuple[str, str], old) -> float:
    sheet, coord = key
    if sheet == CALC:
        return rng.uniform(500, 3000) if coord == 'U3' else rng.uniform(-0.6, 1.2)
    if rng.random() < 0.3:
        return old if isinstance(old, (int, float)) else 0.0
    _, row = split_coord(coord)
    if 1891 <= row <= 1915:                     # 目标伤害减免：让 MIN(x, 1) 的钳位也走到
        return rng.uniform(-0.5, 1.5)
    return rng.uniform(-0.5, 1.2)


def known(g: dict) -> str | None:
    d = g['z'].get('def')
    if d and d['targetDef'] * (1 + d['defRate']) * (1 - d['ignore']) / (800 + d['level'] * 8) + 1 <= 0:
        return '防御分母 ≤ 0'
    return None


def main() -> int:
    ap = argparse.ArgumentParser(description='golden 扰动对拍（TD-03 §10.3）')
    ap.add_argument('xlsx', nargs='?')
    ap.add_argument('--rounds', type=int, default=20)
    ap.add_argument('--seed', type=int, default=1)
    argv = sys.argv[1:]
    if argv[:1] == ['--']:  # pnpm 10 把 `pnpm golden:perturb -- …` 里的 `--` 原样传进来
        argv = argv[1:]
    args = ap.parse_args(argv)
    xlsx = Path(args.xlsx) if args.xlsx else next(iter(sorted((ROOT / 'data' / 'raw').glob('*.xlsx'))), None)
    if xlsx is None:
        print('找不到 xlsx', file=sys.stderr)
        return 2
    t0 = time.time()
    warnings.filterwarnings('ignore')
    print(f'读取 {xlsx.name} …', file=sys.stderr)
    wb = load_workbook(xlsx, data_only=True, read_only=True)
    wbf = load_workbook(xlsx, data_only=False, read_only=True)
    book = {name: golden.sheet_values(wb[name]) for name in golden.SHEETS}
    fbook = {name: golden.sheet_formulas(wbf[name]) for name in (CALC, CFG)}
    base_ev = Evaluator(book, fbook)
    cells = golden.golden_cells(book[CALC], fbook[CALC])
    inputs, live = inputs_of(base_ev, cells)
    print(f'{len(cells)} 个标准答案格；改写 {len(inputs)} 格，现算 {len(live)} 格；{time.time() - t0:.0f} 秒', file=sys.stderr)

    total, equal = 0, 0
    kinds: Counter = Counter()
    worst: dict[str, float] = {}
    samples: list[str] = []
    for rnd in range(args.rounds):
        rng = random.Random(args.seed * 1000 + rnd)
        overrides = {k: draw(rng, k, book[k[0]].get(k[1])) for k in sorted(inputs)}
        ev = Evaluator(book, fbook, live=live, overrides=overrides)
        ev._parsed = base_ev._parsed                      # 解析结果共用
        cls = golden.Classifier(ev)
        for coord, name, br in cells:
            src = fbook[CALC][coord]
            try:
                expected = ev.number(ev.parsed(src), CALC)
                g, why, _ = cls.classify(coord, src, name, br)
            except (XlError, ZeroDivisionError) as e:
                kinds[f'求值出错 {e!r}'] += 1
                continue
            if g is None:
                continue
            total += 1
            got = golden.recompute(g)
            if 'other' in g['z']:
                kinds['有未归类因子'] += 1
            if got == expected:
                equal += 1
                continue
            k = known(g) or '意外'
            kinds[k] += 1
            rel = abs(got - expected) / max(abs(expected), 1)
            tag = f'{k}（差 1）' if abs(got - expected) == 1 else k
            if abs(got - expected) == 1:
                kinds[tag] += 1
            worst[k] = max(worst.get(k, 0.0), rel)
            if k == '意外' and len(samples) < 20:
                samples.append(f'第 {rnd + 1} 轮 {coord} {name}：重算 {got:g}，原公式 {expected:g}')
        print(f'第 {rnd + 1} 轮：累计 {total} 次，全等 {equal}；{time.time() - t0:.0f} 秒', file=sys.stderr)

    print(f'\n{args.rounds} 轮共 {total} 次比较，全等 {equal}，不等 {total - equal}。')
    for k, n in sorted(kinds.items()):
        extra = f'，最大相对差 {worst[k]:.2e}' if k in worst else ''
        print(f'- {k}：{n}{extra}')
    for s in samples:
        print(f'  {s}')
    return 1 if kinds.get('意外') or kinds.get('有未归类因子') else 0


if __name__ == '__main__':
    sys.exit(main())
