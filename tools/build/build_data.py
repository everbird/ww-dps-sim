"""数据构建：xlsx → data/generated/（TD-01）。

用法：
  python3 tools/build/build_data.py [xlsx 路径] [--strict 椿,散华,维里奈]

xlsx 缺省取 data/raw/ 里唯一的 .xlsx。只有 TD-01 §12.4 的阻断项会让退出码非零：
设置区不符合期望、--strict 名单内角色有未解决的伤害判定。zod 校验由 `pnpm check:data` 负责。
"""
from __future__ import annotations

import argparse
import hashlib
import json
import sys
import time
import warnings
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).parent))
from actions import parse_action_sheet  # noqa: E402
from characters import build_characters  # noqa: E402
from dmg import index_dmg, join_block  # noqa: E402
from echoes import build_costs, drop_phantoms, join_echo, parse_echo_sheet  # noqa: E402
from enemies import build_enemies  # noqa: E402
import golden  # noqa: E402
import nanoka  # noqa: E402
from parse import Issues, as_num  # noqa: E402
from report import render, render_echoes, render_golden  # noqa: E402
from weapons import build_weapons  # noqa: E402
from xlformula import Evaluator  # noqa: E402

ROOT = Path(__file__).resolve().parents[2]
EXPECTED_SETTINGS = {'framerate': 60, 'halfFrameThreshold': 0.5, 'coop': False,
                     'limitDodgeBulletTimeCanceling': False, 'weakPointHitting': False}


def merged_map(ws) -> dict:
    """{(行, 列1起): (首行, 末行)}，只记多行合并"""
    mp = {}
    for rng in ws.merged_cells.ranges:
        if rng.max_row > rng.min_row:
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    mp[(r, c)] = (rng.min_row, rng.max_row)
    return mp


def main() -> int:
    ap = argparse.ArgumentParser(description='xlsx → data/generated')
    ap.add_argument('xlsx', nargs='?')
    ap.add_argument('--out', default=str(ROOT / 'data' / 'generated'))
    ap.add_argument('--curated', default=str(ROOT / 'data' / 'curated'))
    ap.add_argument('--strict', default='', help='逗号分隔的角色名：这些角色的伤害判定必须全部连上 dmg')
    ap.add_argument('--nanoka', default='live', help="nanoka 版本（live / latest / 3.7…）；off = 不取（沿用上次的 nanoka.json）")
    argv = sys.argv[1:]
    if argv[:1] == ['--']:  # pnpm 10 把 `pnpm build:data -- …` 里的 `--` 原样传进来
        argv = argv[1:]
    args = ap.parse_args(argv)

    if args.xlsx:
        xlsx = Path(args.xlsx)
    else:
        found = sorted((ROOT / 'data' / 'raw').glob('*.xlsx'))
        if len(found) != 1:
            print(f'data/raw/ 里应当恰好有一个 .xlsx，现在有 {len(found)} 个；也可以直接传路径', file=sys.stderr)
            return 2
        xlsx = found[0]
    t0 = time.time()
    warnings.filterwarnings('ignore')
    print(f'读取 {xlsx.name} …', file=sys.stderr)
    wb = load_workbook(xlsx, data_only=True)                          # 缓存值 + 合并区
    wbf = load_workbook(xlsx, data_only=False, read_only=True)        # 公式
    issues = Issues()
    blocking: list[str] = []

    # §1.8 设置区
    b = wb['base']
    settings = {'framerate': b['B2'].value, 'halfFrameThreshold': b['C2'].value, 'coop': b['B3'].value,
                'limitDodgeBulletTimeCanceling': b['B4'].value, 'weakPointHitting': b['B5'].value}
    for k, v in EXPECTED_SETTINGS.items():
        if settings[k] != v:
            blocking.append(f'设置区 {k} = {settings[k]!r}，期望 {v!r}（TD-01 §1.8）')

    # §3 动作表
    blocks = []
    for sheet, gender in (('角色-女', '女'), ('角色-男', '男')):
        ws = wb[sheet]
        values = list(ws.iter_rows(min_row=1, max_col=46, values_only=True))
        formulas = list(wbf[sheet].iter_rows(min_row=1, max_col=46, values_only=True))
        blocks += parse_action_sheet(sheet, gender, values, formulas, merged_map(ws), issues)

    # §4 dmg 连接
    dmg = index_dmg(list(wb['dmg'].iter_rows(min_row=3, max_col=110, values_only=True)), issues)
    join_path = Path(args.curated) / 'dmg-join.json'
    dmg_join = json.loads(join_path.read_text('utf-8')) if join_path.exists() else {}
    dmg_join = {k: v for k, v in dmg_join.items() if not k.startswith('$')}
    gaps, cross = [], []
    for blk in blocks:
        join_block(blk, dmg, dmg_join, issues, gaps, cross)

    # §5 / §6 角色
    base_rows = list(b.iter_rows(min_row=56, max_col=180, values_only=True))
    index_rows = list(wb['索引'].iter_rows(min_row=1, max_row=47, max_col=23, values_only=True))
    characters = build_characters(base_rows, index_rows, {blk['key']: blk for blk in blocks}, issues)

    # §7 武器、§9 敌人
    weapons = build_weapons(list(wb['weapon'].iter_rows(min_row=4, max_col=61, values_only=True)), base_rows, issues)
    enemies = build_enemies(list(wb['敌对属性列表'].iter_rows(min_row=2, max_col=23, values_only=True)), issues)

    # §8 声骸（倍率在 nanoka 之后连）
    ws = wb['声骸']
    costs = build_costs(list(wb['索引'].iter_rows(min_row=48, max_row=58, max_col=23, values_only=True)))
    echoes, phantoms = drop_phantoms(parse_echo_sheet(
        list(ws.iter_rows(min_row=1, max_col=34, values_only=True)),
        list(wbf['声骸'].iter_rows(min_row=1, max_col=34, values_only=True)), merged_map(ws), costs, issues))

    # nanoka：技能冷却、技能文本与声骸数据（m0-confirm §5）。取不到只提示，不阻断；声骸倍率沿用上次 nanoka.json 里的
    out = Path(args.out)
    nk, nk_note = None, ''
    if args.nanoka != 'off':
        try:
            nk, notes = nanoka.build(list(characters), args.nanoka, nanoka.load_name_map(Path(args.curated)),
                                     [e['key'] for e in echoes], nanoka.load_echo_name_map(Path(args.curated)))
            nk_note = (f"\n## nanoka\n\n版本 {nk['version']}，{len(nk['characters'])} 个角色、{len(nk['echoes'])} 个声骸 → `nanoka.json`。\n"
                       + ''.join(f'\n- {n}' for n in notes) + ('\n' if notes else ''))
        except Exception as e:  # noqa: BLE001
            print(f'nanoka 没取到（{e}）；沿用上次的 nanoka.json', file=sys.stderr)
            nk_note = f'\n## nanoka\n\n没取到：{e}\n'
    nk_echoes = (nk or _previous_nanoka(out)).get('echoes', {})
    echo_gaps: list = []
    for e in echoes:
        join_echo(e, dmg, dmg_join, nk_echoes.get(e['key']), issues, echo_gaps)

    # §5.1 公式定义原文
    formula_ref = {}
    e4 = b['E4'].value
    if isinstance(e4, str) and '=' in e4:
        name, text = e4.split('=', 1)
        formula_ref[name.strip()] = text.strip()
    for row in range(5, 11):
        name, text = b[f'D{row}'].value, b[f'E{row}'].value
        if isinstance(name, str) and text is not None:
            formula_ref[name.strip()] = str(text).strip()

    # golden：标准答案与逐格乘区（TD-01 §11.1、TD-03 §10）
    book = {name: golden.sheet_values(wb[name]) for name in golden.SHEETS}
    fbook = {name: golden.sheet_formulas(wbf[name]) for name in (golden.CALC, golden.CFG)}
    golden_damage = golden.golden_damage(book[golden.CALC], book[golden.CFG].get('B3'), golden.dmg_by_calc(book['dmg']))
    golden_zones, golden_info = golden.golden_zones(Evaluator(book, fbook))

    # --strict
    strict_names = [s.strip() for s in args.strict.replace('，', ',').split(',') if s.strip()]
    strict = {}
    for name in strict_names:
        if not any(blk['key'] == name for blk in blocks):
            blocking.append(f'--strict 里的 {name} 在动作表里找不到')
            continue
        strict[name] = [g for g in gaps if g['block'] == name]
        if strict[name]:
            blocking.append(f'严格名单：{name} 有 {len(strict[name])} 个伤害判定没连上 dmg（见下节）')

    # 写文件
    (out / 'actions').mkdir(parents=True, exist_ok=True)
    for old in (out / 'actions').glob('*.json'):
        old.unlink()
    for blk in blocks:
        _dump(out / 'actions' / f"{blk['key']}.json", blk)
    _dump(out / 'characters.json', characters)
    _dump(out / 'formula-ref.json', formula_ref)
    _dump(out / 'weapons.json', weapons)
    _dump(out / 'echoes.json', echoes)
    _dump_lines(out / 'enemies.json', enemies)
    if nk is not None:
        _dump(out / 'nanoka.json', nk)
    (out / 'fixtures').mkdir(exist_ok=True)
    _dump(out / 'fixtures' / 'golden-damage.json', golden.clean(golden_damage))
    _dump_lines(out / 'fixtures' / 'golden-zones.json', golden.clean(golden_zones))

    rows = [r for blk in blocks for g in blk['groups'] for r in g['rows']]
    kinds = Counter(r['kind'] for r in rows)
    hit = [r for r in rows if r['kind'] == 'hit']
    flags = Counter(f for r in rows for f in r['flags'])
    counts = {
        '块（角色 + 通用）': len(blocks), '动作组': sum(len(blk['groups']) for blk in blocks),
        '数据行': len(rows), '判定行': kinds['hit'], '资源行': kinds['gain'], '膨胀行': kinds['dilation'],
        '标记行': kinds['marker'], '事件生成的判定': sum(1 for r in hit if r['eventSpawned']),
        '忽略的无名行': sum(blk['ignoredRows'] for blk in blocks),
        '连上 dmg（直接 / 别名）': f"{sum(1 for r in hit if r.get('dmg', {}).get('via') == 'direct')} / "
                              f"{sum(1 for r in hit if r.get('dmg', {}).get('via') == 'alias')}",
        'dmgNoConfig': flags['dmgNoConfig'], '伤害判定没连上（noDmg）': flags['noDmg'],
        '出生帧': sum(1 for r in rows if r['birthFrame'] is not None),
        '角色（characters.json）': len(characters),
        '武器（weapons.json）': len(weapons),
        '声骸（echoes.json）': len(echoes),
        '敌人（enemies.json）': len(enemies),
        'golden 条目（带 dmgKey）': f"{len(golden_damage['entries'])}（{sum(1 for e in golden_damage['entries'] if 'dmgKey' in e)}）",
        'golden 乘区格': len(golden_zones),
    }
    meta = {
        'xlsxFile': xlsx.name, 'xlsxSha256': hashlib.sha256(xlsx.read_bytes()).hexdigest(),
        'resourceVersion': _resource_version(wb), 'builtAt': datetime.now().astimezone().isoformat(timespec='seconds'),
        'settings': settings,
        'counts': {'blocks': len(blocks), 'groups': counts['动作组'], 'rows': len(rows), 'hitRows': kinds['hit'],
                   'joined': sum(1 for r in hit if 'dmg' in r), 'noDmg': flags['noDmg'],
                   'characters': len(characters), 'weapons': len(weapons), 'echoes': len(echoes), 'enemies': len(enemies),
                   'goldenEntries': len(golden_damage['entries']),
                   'goldenZones': len(golden_zones)},
    }
    _dump(out / 'meta.json', meta)
    report = render(meta, counts, blocking, strict, gaps, cross, issues, blocks, flags)
    report += render_echoes(echoes, echo_gaps, bool(nk_echoes), phantoms)
    report += render_golden(golden_damage, golden_zones, golden_info)
    report += nk_note
    (out / 'build-report.md').write_text(report, 'utf-8')

    print(f"完成：{len(blocks)} 块、{counts['动作组']} 组、{len(rows)} 行；连上 {meta['counts']['joined']}，"
          f"noDmg {flags['noDmg']}；{time.time() - t0:.0f} 秒。报告：{out / 'build-report.md'}", file=sys.stderr)
    for msg in blocking:
        print(f'阻断：{msg}', file=sys.stderr)
    return 1 if blocking else 0


def _previous_nanoka(out: Path) -> dict:
    """--nanoka off 或没取到时，沿用上次构建的 nanoka.json（声骸倍率要用）"""
    p = out / 'nanoka.json'
    return json.loads(p.read_text('utf-8')) if p.exists() else {}


def _resource_version(wb) -> str | None:
    v = wb['索引']['T60'].value
    if isinstance(v, str) and 'Version' in v:
        return v.split('Version', 1)[1].strip()
    return None


def _dump(path: Path, obj) -> None:
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=1, default=as_num) + '\n', 'utf-8')


def _dump_lines(path: Path, items: list) -> None:
    """数组每项一行（条目多、每条结构相同时更好读、diff 也更干净）"""
    body = ',\n'.join(json.dumps(x, ensure_ascii=False, default=as_num) for x in items)
    path.write_text(f'[\n{body}\n]\n', 'utf-8')


if __name__ == '__main__':
    sys.exit(main())
